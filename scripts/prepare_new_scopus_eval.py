"""Freeze unseen August Scopus papers for an independent Astra evaluation."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from prepare_astra_reference import normalized_title, write_once


AUGUST = Path("C:/Users/aumki/Downloads/scopus_full_export_2026-08-16-10-05-15.xlsx")
SEPTEMBER = Path("C:/Users/aumki/Downloads/scopus_full_export_2026-09-15-15-07-38_v003.xlsx")
DESTINATION = Path("evaluation/new-scopus-v1")
SEED = "new-scopus-v1-2026-09-23"


def read_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = workbook["Documents"].iter_rows(values_only=True)
        headers = next(iterator)
        required = {"scopus_id", "title", "abstract", "doi", "authkeywords", "scopus_link"}
        if not required <= set(headers):
            raise ValueError(f"Missing Scopus columns in {path}")
        rows = []
        for row_number, cells in enumerate(iterator, start=2):
            row = dict(zip(headers, cells))
            if not row.get("scopus_id"):
                continue
            keywords = row.get("authkeywords")
            if isinstance(keywords, str):
                try:
                    parsed = json.loads(keywords)
                    keywords = parsed if isinstance(parsed, list) else keywords
                except json.JSONDecodeError:
                    pass
            rows.append({"row": row_number, "scopus_id": str(row["scopus_id"]),
                         "title": str(row.get("title") or "").strip(),
                         "abstract": str(row.get("abstract") or "").strip(),
                         "authkeywords": keywords, "doi": str(row.get("doi") or "").strip(),
                         "scopus_link": row.get("scopus_link")})
        return rows
    finally:
        workbook.close()


def rank(key: str) -> str:
    return hashlib.sha256(f"{SEED}:{key}".encode("utf-8")).hexdigest()


def main() -> None:
    old, incoming = read_rows(SEPTEMBER), read_rows(AUGUST)
    old_ids = {row["scopus_id"] for row in old}
    old_titles = {normalized_title(row["title"]) for row in old if row["title"]}
    old_dois = {row["doi"].casefold() for row in old if row["doi"]}
    candidates = [row for row in incoming if row["scopus_id"] not in old_ids
                  and normalized_title(row["title"]) not in old_titles
                  and (not row["doi"] or row["doi"].casefold() not in old_dois)
                  and (row["abstract"] or row["authkeywords"])]
    # Keep one record per paper, including alternate Scopus IDs for the same DOI/title.
    seen_titles, seen_dois = set(), set()
    unique = []
    for row in sorted(candidates, key=lambda row: row["scopus_id"]):
        title, doi = normalized_title(row["title"]), row["doi"].casefold()
        if title in seen_titles or (doi and doi in seen_dois):
            continue
        unique.append(row)
        seen_titles.add(title)
        if doi:
            seen_dois.add(doi)
    random_sample = sorted(unique, key=lambda row: rank(row["scopus_id"]))[:100]
    random_ids = {row["scopus_id"] for row in random_sample}
    challenge = []
    for row in sorted(unique, key=lambda row: rank("challenge:" + row["scopus_id"])):
        if row["scopus_id"] in random_ids:
            continue
        text = (row["title"] + " " + row["abstract"][:1500]).casefold()
        if (re.search(r"\b(ai|machine learning|deep learning|neural|transformer|bert|cnn)\b", text)
                and re.search(r"\b(novel|propos\w*|architecture|method|framework|model)\b", text)
                and re.search(r"\b(health|medical|tumou?r|image|disease|classification|predict\w*|language|sentiment)\b", text)):
            challenge.append(row)
        if len(challenge) == 30:
            break
    if len(random_sample) != 100 or len(challenge) != 30:
        raise ValueError(f"Insufficient independent records: random={len(random_sample)}, challenge={len(challenge)}")
    selected = sorted(random_sample + challenge, key=lambda row: row["row"])
    manifest = {"schema_version": "new-scopus-v1", "selection_seed": SEED,
                "source_sha256": hashlib.sha256(AUGUST.read_bytes()).hexdigest(),
                "excluded_source_sha256": hashlib.sha256(SEPTEMBER.read_bytes()).hexdigest(),
                "selection_policy": "Exclude matching Scopus IDs, normalized titles, or DOIs; deduplicate August; select 100 hash-ranked random records and 30 distinct AI-boundary challenge records.",
                "eligible_unique": len(unique), "random_ids": sorted(random_ids),
                "challenge_ids": sorted(row["scopus_id"] for row in challenge),
                "total_selected": len(selected)}
    DESTINATION.mkdir(parents=True, exist_ok=True)
    write_once(DESTINATION / "manifest.json", (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    # Raw paper text is kept in ignored outputs; the tracked manifest freezes its IDs and source hash.
    output = Path("outputs/new-scopus-v1/source_records.json")
    output.parent.mkdir(parents=True, exist_ok=True)
    write_once(output, (json.dumps(selected, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    print(json.dumps({"august_rows": len(incoming), "overlapping_ids": len(old_ids & {r["scopus_id"] for r in incoming}),
                      "eligible_unique": len(unique), "random": 100, "challenge": 30}))


if __name__ == "__main__":
    main()
