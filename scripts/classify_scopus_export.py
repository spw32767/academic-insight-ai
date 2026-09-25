"""Classify Scopus export rows locally; save resumable JSON for workbook review.

This script reads the workbook but never edits it. The reviewed JSON is written
back to a *copy* of the workbook in a separate spreadsheet-authoring step.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import time
from pathlib import Path

from openpyxl import load_workbook

from academic_insight_ai.apis.common import build_provider
from academic_insight_ai.tasks.article_classification.fewshot_examples import BOUNDARY_EXAMPLE_IDS, EXAMPLE_IDS
from academic_insight_ai.tasks.article_classification.service import (
    AGREEMENT_CLASSIFIER_VERSION, ASTRA_CANDIDATE_VERSION, BOUNDARY_CANDIDATE_VERSION,
    CLASSIFIER_VERSION, FEWSHOT_CANDIDATE_VERSION,
    LEGACY_CLASSIFIER_VERSION,
    ClassificationRequest, classify,
)
from paper_categories import CATEGORIES


TAXONOMY_VERSION = "2026-09-21-seven-category"


def read_records(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Documents"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        required = {"scopus_id", "scopus_link", "title", "abstract", "authkeywords", "category", "classification_confidence"}
        missing = required - set(headers)
        if missing:
            raise ValueError(f"Missing expected columns: {sorted(missing)}")
        records = []
        for row_number, cells in enumerate(rows, start=2):
            row = dict(zip(headers, cells))
            if not row.get("scopus_id"):
                continue
            if row.get("category") or row.get("classification_confidence"):
                raise ValueError(f"Row {row_number} already has a classification; refusing to overwrite it")
            keywords = row.get("authkeywords")
            if isinstance(keywords, str):
                try:
                    parsed = json.loads(keywords)
                    keywords = parsed if isinstance(parsed, list) else keywords
                except json.JSONDecodeError:
                    pass
            records.append({
                "row": row_number,
                "scopus_id": row["scopus_id"],
                "scopus_link": row["scopus_link"],
                "title": row["title"],
                "abstract": row["abstract"],
                "authkeywords": keywords,
            })
        return records
    finally:
        workbook.close()


def save_checkpoint(path: Path, report: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--checkpoint", type=Path, default=Path("outputs/evidence-agreement/checkpoint.json"))
    parser.add_argument("--model", default="qwen3-4b")
    parser.add_argument("--limit", type=int, help="Classify at most this many new rows for a pilot run")
    parser.add_argument("--review-preface", action="store_true", help="Recheck substantive abstracts marked Preface")
    parser.add_argument("--verification-mode", choices=["legacy", "single", "evidence_agreement", "astra_candidate", "fewshot_candidate", "boundary_candidate"],
                        default="evidence_agreement", help="Use agreement mode for the conservative experiment")
    parser.add_argument("--split-manifest", type=Path, help="Frozen split manifest for a partitioned evaluation")
    parser.add_argument("--partition", choices=["development", "holdout"], help="Partition to classify")
    parser.add_argument("--pilot-manifest", type=Path, help="Run only the disjoint few-shot pilot IDs")
    args = parser.parse_args()

    source_hash = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    records = read_records(args.workbook)
    if bool(args.split_manifest) != bool(args.partition):
        raise ValueError("--split-manifest and --partition must be supplied together")
    split_hash = None
    if args.split_manifest:
        split_bytes = args.split_manifest.read_bytes()
        split_hash = hashlib.sha256(split_bytes).hexdigest()
        split = json.loads(split_bytes)
        if split["source_sha256"] != source_hash:
            raise ValueError("Split manifest belongs to another workbook")
        selected_ids = set(split[f"{args.partition}_ids"])
        if not selected_ids <= {str(record["scopus_id"]) for record in records}:
            raise ValueError("Split manifest contains IDs absent from workbook")
        records = [record for record in records if str(record["scopus_id"]) in selected_ids]
    pilot_hash = None
    if args.pilot_manifest:
        if args.partition != "development" or args.verification_mode != "fewshot_candidate":
            raise ValueError("Few-shot pilot requires development partition and fewshot_candidate mode")
        pilot_bytes = args.pilot_manifest.read_bytes()
        pilot_hash = hashlib.sha256(pilot_bytes).hexdigest()
        pilot = json.loads(pilot_bytes)
        if pilot["source_sha256"] != source_hash or set(pilot["example_ids"]) != EXAMPLE_IDS:
            raise ValueError("Few-shot pilot does not match workbook or prompt examples")
        pilot_ids = set(pilot["pilot_ids"])
        if pilot_ids & EXAMPLE_IDS or not pilot_ids <= {str(record["scopus_id"]) for record in records}:
            raise ValueError("Few-shot pilot overlaps examples or lies outside development partition")
        records = [record for record in records if str(record["scopus_id"]) in pilot_ids]
    if args.verification_mode in ("fewshot_candidate", "boundary_candidate") and args.partition == "development":
        examples = BOUNDARY_EXAMPLE_IDS if args.verification_mode == "boundary_candidate" else EXAMPLE_IDS
        records = [record for record in records if str(record["scopus_id"]) not in examples]
    if len({str(record["scopus_id"]) for record in records}) != len(records):
        raise ValueError("Duplicate scopus_id values in workbook")
    if args.checkpoint.exists():
        report = json.loads(args.checkpoint.read_text(encoding="utf-8"))
        expected_version = (AGREEMENT_CLASSIFIER_VERSION if args.verification_mode == "evidence_agreement"
                            else CLASSIFIER_VERSION if args.verification_mode == "single"
                            else ASTRA_CANDIDATE_VERSION if args.verification_mode == "astra_candidate"
                            else FEWSHOT_CANDIDATE_VERSION if args.verification_mode == "fewshot_candidate"
                            else BOUNDARY_CANDIDATE_VERSION if args.verification_mode == "boundary_candidate"
                            else LEGACY_CLASSIFIER_VERSION)
        if (report["source_sha256"] != source_hash or report["model"] != args.model
                or report.get("classifier_version") != expected_version
                or report.get("verification_mode", "evidence_agreement") != args.verification_mode
                or report.get("taxonomy_version") != TAXONOMY_VERSION
                or report.get("split_sha256") != split_hash or report.get("partition") != args.partition
                or report.get("pilot_sha256") != pilot_hash):
            raise ValueError("Checkpoint belongs to another workbook, model, taxonomy, or classifier version")
    else:
        report = {
            "source_file": str(args.workbook.resolve()),
            "source_sha256": source_hash,
            "model": args.model,
            "taxonomy_version": TAXONOMY_VERSION,
            "classifier_version": (AGREEMENT_CLASSIFIER_VERSION if args.verification_mode == "evidence_agreement"
                                   else CLASSIFIER_VERSION if args.verification_mode == "single"
                                   else ASTRA_CANDIDATE_VERSION if args.verification_mode == "astra_candidate"
                                   else FEWSHOT_CANDIDATE_VERSION if args.verification_mode == "fewshot_candidate"
                                   else BOUNDARY_CANDIDATE_VERSION if args.verification_mode == "boundary_candidate"
                                   else LEGACY_CLASSIFIER_VERSION),
            "verification_mode": args.verification_mode,
            "split_sha256": split_hash,
            "partition": args.partition,
            "pilot_sha256": pilot_hash,
            "total_records": len(records),
            "results": {},
        }

    provider, model_name = build_provider(args.model)
    results = report["results"]
    completed = 0
    for record in records:
        key = str(record["scopus_id"])
        if args.review_preface:
            previous = results.get(key)
            if not previous or previous["confidence"] != "Preface" or previous.get("reviewed_substantive_preface"):
                continue
            if str(record["title"] or "").lower().startswith(("preface", "editorial", "correction", "erratum", "retraction notice")):
                continue
            if len(str(record["abstract"] or "").strip()) <= 100:
                continue
        elif key in results:
            continue
        if args.limit is not None and completed >= args.limit:
            break
        request = ClassificationRequest(
            paper_id=key,
            title=str(record["title"] or ""),
            abstract=record["abstract"],
            authkeywords=record["authkeywords"],
            categories=CATEGORIES,
            taxonomy_version=TAXONOMY_VERSION,
            verification_mode=args.verification_mode,
        )
        started = time.perf_counter()
        for attempt in range(3):
            try:
                classification = classify(request, provider, model_name)
                break
            except Exception:
                if attempt == 2:
                    raise
                time.sleep(2 * (attempt + 1))
        updated = {
            "row": record["row"],
            "scopus_id": key,
            "title": record["title"],
            "scopus_link": record["scopus_link"],
            "category_code": classification.primary_category_code,
            "confidence": classification.confidence,
            "model_reported_confidence": classification.model_reported_confidence,
            "confidence_source": classification.confidence_source,
            "reason": classification.reason,
            "verification_status": classification.verification_status,
            "primary_contribution": classification.primary_contribution,
            "evidence_quote": classification.evidence_quote,
            "evidence_source": classification.evidence_source,
            "candidate_category_codes": classification.candidate_category_codes,
            "assessments": [assessment.model_dump() for assessment in classification.assessments],
            "elapsed_seconds": round(time.perf_counter() - started, 2),
        }
        if args.review_preface:
            updated["reviewed_substantive_preface"] = True
            updated["previous_result"] = previous
        results[key] = updated
        save_checkpoint(args.checkpoint, report)
        completed += 1
        print(f"{len(results)}/{len(records)} row {record['row']}: {classification.primary_category_code or 'NULL'} / {classification.confidence}", flush=True)
    print(f"Saved {len(results)}/{len(records)} classifications to {args.checkpoint.resolve()}", flush=True)


if __name__ == "__main__":
    main()
