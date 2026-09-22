"""Verify that only the category and confidence cells changed in an XLSX copy."""

from __future__ import annotations

import argparse
from collections import Counter

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("classified")
    args = parser.parse_args()

    source = load_workbook(args.source)
    classified = load_workbook(args.classified)
    assert source.sheetnames == classified.sheetnames == ["Documents"]
    before = source["Documents"]
    after = classified["Documents"]
    assert before.max_row == after.max_row == 164
    assert before.max_column == after.max_column == 38
    assert before.freeze_panes == after.freeze_panes
    assert list(before.merged_cells.ranges) == list(after.merged_cells.ranges)

    confidence_counts: Counter[str] = Counter()
    for row in range(1, 165):
        for col in range(1, 39):
            left, right = before.cell(row, col), after.cell(row, col)
            if row == 1 or col not in (37, 38):
                assert left.value == right.value, f"Changed data at {left.coordinate}"
            assert left.number_format == right.number_format, f"Changed number format at {left.coordinate}"
        if row > 1:
            category = after.cell(row, 37).value
            confidence = after.cell(row, 38).value
            assert confidence in {"High", "Medium", "Low", "Preface"}, (row, confidence)
            assert (category is None) == (confidence == "Preface"), (row, category, confidence)
            if category is not None:
                assert category in range(1, 8), (row, category)
            confidence_counts[confidence] += 1

    assert sum(confidence_counts.values()) == 163
    print(f"Verified 163 rows; only AK:AL values changed. Confidence counts: {dict(confidence_counts)}")


if __name__ == "__main__":
    main()
